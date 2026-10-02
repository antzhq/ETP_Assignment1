"""
Hybrid CSP + PVT plant sizing for Santiago (Assignment 1, Question 5, with Q6 and Q7 graphs).

This script implements the mathematical model in "../2_Math_Model/Q5_Mathematical_Model.pdf" one-to-one.
Every variable carries the same name as its symbol in that document, e.g.
    Yield_{PVT heat,hr,day}  ->  yield_PVT_heat[day][hr]
    Area_{CSP}               ->  area_CSP

What it does, in order:
    1. Reads the hourly data (irradiance, ambient temperature, demand) from the Excel file.
    2. Calculates the derived parameters (efficiency chains, yield per m2 for every hour).
    3. Solves the model: finds the cheapest Count_heliostats and Count_PVT_panels
       that satisfy constraints C1-C5 on both days.
    4. Prints the results and saves three figures to ../4_Visualizations:
         Q5_Design_Space_Cost_Map.png     cost of every candidate design, feasible region, optimum
         Q6_Hourly_Supply_vs_Demand.png   hourly supply vs demand, both days, electricity and heat
         Q7_Mirror_Degradation.png        plant output with mirror reflectivity 0.95 vs 0.85
    5. Exports the data behind the figures to ../5_Data/Results_Plant_Sizing.xlsx (live formulas).

How to run (from anywhere; paths are resolved relative to this file):
    pip install numpy matplotlib openpyxl
    python plant_sizing.py                      (reads ../5_Data/Input_Data_Assignment_1.xlsx)
    python plant_sizing.py "path/to/Data Assignment 1.xlsx"
"""

import math
import sys

import matplotlib.pyplot as plt
import numpy as np
import openpyxl

# =============================================================================
# 1. SETS
# =============================================================================
days = ["June", "December"]  # Days
hrs = list(range(24))         # Hrs = {0, 1, ..., 23}

# =============================================================================
# 2. PARAMETERS (given in the assignment brief)
# =============================================================================
# --- CSP ---
eff_reflectivity = 0.95       # heliostat mirror reflectivity (Q7 compares 0.85)
eff_cosine = 0.86             # average cosine efficiency
eff_shading_blocking = 0.89   # shading and blocking efficiency
eff_atmospheric = 0.97        # atmospheric attenuation efficiency
eff_spillage = 0.92           # spillage efficiency
eff_receiver = 0.80           # receiver absorption efficiency
eff_rankine = 0.42            # Rankine cycle thermal efficiency
area_per_heliostat = 220.0    # m2 of mirror per heliostat
cost_CSP = 80.0               # EUR per m2 of heliostat mirror

# --- PVT ---
eff_PVT_elec = 0.16           # PVT electrical efficiency
frac_absorbed = 0.86          # effective optical fraction, textbook (tau*alpha)
coef_heat_loss = 12.0         # overall heat loss coefficient U_L, W/(m2 K)
factor_heat_removal = 0.65    # collector heat removal factor F_R
T_water_in = 35.0             # PVT inlet water temperature, degC
flow_rate_per_PVT_panel = 0.2 # kg/s of water through one 2 m x 2 m panel (used for the outlet temperature)
cp_water = 4186.0             # J/(kg K), specific heat of water (standard value, not given in the brief)
area_per_PVT_panel = 4.0      # m2 per PVT panel (2 m x 2 m)
cost_PVT = 150.0              # EUR per m2 of PVT panel

# --- Plant ---
area_max = 450_000.0          # maximum available area, m2
share_CSP_min = 0.60          # minimum share of electricity from CSP
GJ_per_W_hr = 3600 / 1e9      # energy of 1 W sustained for 1 hour, in GJ (= 3.6e-6)


# =============================================================================
# 2b. DATA (hourly, read from the Excel file)
# =============================================================================
def read_data(excel_path):
    """
    Reads the four hourly data series for both days from the assignment's Excel file.

    Sheet 'Supply' : irradiance [W/m2] and ambient temperature [degC]
    Sheet 'Demand' : heat and electricity demand per hour [GJ]
    Data rows are rows 4 to 27 (hours 0 to 23) in both sheets.

    Returns four dictionaries indexed as  variable[day][hr].
    """
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    sheet_supply = workbook["Supply"]
    sheet_demand = workbook["Demand"]

    # Column letter of each series, per day, as laid out in the Excel file
    columns = {
        "June":     {"irradiance": "B", "T_ambient": "C", "demand_heat": "B", "demand_elec": "D"},
        "December": {"irradiance": "E", "T_ambient": "F", "demand_heat": "K", "demand_elec": "M"},
    }
    first_data_row = 4

    irradiance, T_ambient, demand_heat, demand_elec = {}, {}, {}, {}
    for day in days:
        col = columns[day]
        rows = [first_data_row + hr for hr in hrs]
        irradiance[day] = np.array([float(sheet_supply[f"{col['irradiance']}{r}"].value) for r in rows])
        T_ambient[day] = np.array([float(sheet_supply[f"{col['T_ambient']}{r}"].value) for r in rows])
        demand_heat[day] = np.array([float(sheet_demand[f"{col['demand_heat']}{r}"].value) for r in rows])
        demand_elec[day] = np.array([float(sheet_demand[f"{col['demand_elec']}{r}"].value) for r in rows])
    return irradiance, T_ambient, demand_heat, demand_elec


# =============================================================================
# 3. DERIVED PARAMETERS
# =============================================================================
def efficiency_chain(reflectivity):
    """Eff_optical and Eff_CSP (sunlight on mirrors -> electricity) for a given mirror reflectivity."""
    eff_optical = reflectivity * eff_cosine * eff_shading_blocking * eff_atmospheric * eff_spillage
    eff_CSP = eff_optical * eff_receiver * eff_rankine
    return eff_optical, eff_CSP


def yields_per_m2(irradiance, T_ambient, eff_CSP):
    """
    Yield per m2 for every hour and day, in W/m2.

    yield_CSP_elec : electricity from 1 m2 of heliostat mirror
    yield_PVT_elec : electricity from 1 m2 of PVT panel
    yield_PVT_heat : useful heat from 1 m2 of PVT panel, heat-removal-factor form of the
                     flat plate collector (Lecture 2):
                         Q_u / A_c = F_R * [ (tau*alpha) * I_T - U * (T_in - T_a) ]
                     The heat loss is evaluated at the known inlet water temperature and F_R
                     corrects for the plate being warmer than the inlet. As in the lecture's
                     low-temperature form, radiation losses are not modelled separately.
                     Clipped at 0: when losses exceed gains the pump stops (no negative heat).
    """
    yield_CSP_elec, yield_PVT_elec, yield_PVT_heat = {}, {}, {}
    for day in days:
        yield_CSP_elec[day] = eff_CSP * irradiance[day]
        yield_PVT_elec[day] = eff_PVT_elec * irradiance[day]
        heat_absorbed = frac_absorbed * irradiance[day]
        heat_lost = coef_heat_loss * (T_water_in - T_ambient[day])
        yield_PVT_heat[day] = np.maximum(0.0, factor_heat_removal * (heat_absorbed - heat_lost))
    return yield_CSP_elec, yield_PVT_elec, yield_PVT_heat


def outlet_temperature(yield_PVT_heat):
    """
    PVT water outlet temperature [degC] for every hour, from the fluid's energy balance
    Q_u = m_dot * c_p * (T_out - T_in)  ->  T_out = T_in + Q_u_per_panel / (m_dot * c_p).
    """
    return {day: T_water_in + yield_PVT_heat[day] * area_per_PVT_panel / (flow_rate_per_PVT_panel * cp_water)
            for day in days}


def sun_hrs_mask(irradiance):
    """SunHrs_day as a True/False mask over the 24 hours: True where irradiance > 0."""
    return {day: irradiance[day] > 0 for day in days}


# =============================================================================
# 5. MODEL: supply, objective, constraints
# =============================================================================
def supply(area_CSP, area_PVT, yield_CSP_elec, yield_PVT_elec, yield_PVT_heat):
    """
    Supply in every hour, in GJ (energy delivered during that hour).
    Area [m2] x Yield [W/m2] = power [W];  x GJ_per_W_hr = energy in that hour [GJ].
    """
    supply_CSP_elec = {d: GJ_per_W_hr * area_CSP * yield_CSP_elec[d] for d in days}
    supply_PVT_elec = {d: GJ_per_W_hr * area_PVT * yield_PVT_elec[d] for d in days}
    supply_PVT_heat = {d: GJ_per_W_hr * area_PVT * yield_PVT_heat[d] for d in days}
    return supply_CSP_elec, supply_PVT_elec, supply_PVT_heat


def cost_total(area_CSP, area_PVT):
    """Objective function: total investment cost in EUR."""
    return cost_CSP * area_CSP + cost_PVT * area_PVT


def check_constraints(area_CSP, area_PVT, yields, sun_hrs, demand_heat, demand_elec):
    """
    Evaluates C1-C4 for a candidate design. Works on single numbers or on whole numpy grids.
    Returns a dictionary {constraint name: True/False (or array of them)}.
    """
    yield_CSP_elec, yield_PVT_elec, yield_PVT_heat = yields
    satisfied = {}
    for day in days:
        sun = sun_hrs[day]
        # Sums over the sunlit hours, per m2 of collector [GJ/m2] and of demand [GJ]
        energy_per_area_CSP_elec = GJ_per_W_hr * yield_CSP_elec[day][sun].sum()
        energy_per_area_PVT_elec = GJ_per_W_hr * yield_PVT_elec[day][sun].sum()
        energy_per_area_PVT_heat = GJ_per_W_hr * yield_PVT_heat[day][sun].sum()
        demand_sun_heat = demand_heat[day][sun].sum()
        demand_sun_elec = demand_elec[day][sun].sum()

        supply_sun_CSP_elec = area_CSP * energy_per_area_CSP_elec
        supply_sun_PVT_elec = area_PVT * energy_per_area_PVT_elec
        supply_sun_PVT_heat = area_PVT * energy_per_area_PVT_heat

        # C1  heat supplied over sunlit hours >= heat demanded over sunlit hours
        satisfied[f"C1 heat ({day})"] = supply_sun_PVT_heat >= demand_sun_heat
        # C2  electricity supplied >= electricity demanded (over sunlit hours)
        satisfied[f"C2 electricity ({day})"] = supply_sun_CSP_elec + supply_sun_PVT_elec >= demand_sun_elec
        # C3  CSP electricity >= 60% of all electricity produced
        satisfied[f"C3 CSP share ({day})"] = (
            supply_sun_CSP_elec >= share_CSP_min * (supply_sun_CSP_elec + supply_sun_PVT_elec)
        )
    # C4  total collector area within the site limit
    satisfied["C4 area"] = area_CSP + area_PVT <= area_max
    return satisfied


def solve(yields, sun_hrs, demand_heat, demand_elec):
    """
    Finds the cheapest integer design (C5) satisfying C1-C4 on both days. Exact method:

    For every possible Count_PVT_panels (0 ... area_max / 4), compute the SMALLEST
    Count_heliostats that satisfies C2 and C3 on both days. Fewer heliostats would
    violate one of them, and more would only add cost, so this is the best heliostat
    count for that PVT count. Then keep the candidates that also satisfy C1 and C4,
    and pick the cheapest. With only two variables this enumeration is fast and
    guaranteed to find the true integer optimum (no solver needed).
    """
    yield_CSP_elec, yield_PVT_elec, yield_PVT_heat = yields
    count_PVT_panels = np.arange(0, int(area_max // area_per_PVT_panel) + 1)
    area_PVT = count_PVT_panels * area_per_PVT_panel

    area_CSP_needed = np.zeros_like(area_PVT)
    for day in days:
        sun = sun_hrs[day]
        energy_per_area_CSP_elec = GJ_per_W_hr * yield_CSP_elec[day][sun].sum()
        energy_per_area_PVT_elec = GJ_per_W_hr * yield_PVT_elec[day][sun].sum()
        demand_sun_elec = demand_elec[day][sun].sum()
        # From C2:  area_CSP >= (demand - PVT electricity) / CSP electricity per m2
        area_CSP_for_C2 = (demand_sun_elec - area_PVT * energy_per_area_PVT_elec) / energy_per_area_CSP_elec
        # From C3 (rearranged):  area_CSP >= share/(1-share) * PVT elec per m2 / CSP elec per m2 * area_PVT
        area_CSP_for_C3 = (share_CSP_min / (1 - share_CSP_min)) * (
            energy_per_area_PVT_elec / energy_per_area_CSP_elec) * area_PVT
        area_CSP_needed = np.maximum.reduce([area_CSP_needed, area_CSP_for_C2, area_CSP_for_C3])

    # Round UP to whole heliostats (C5). The tiny tolerance avoids 1000.0000001 -> 1001.
    count_heliostats = np.ceil(area_CSP_needed / area_per_heliostat - 1e-9).astype(int)
    area_CSP = count_heliostats * area_per_heliostat

    satisfied = check_constraints(area_CSP, area_PVT, yields, sun_hrs, demand_heat, demand_elec)
    feasible = np.logical_and.reduce(list(satisfied.values()))
    if not feasible.any():
        raise RuntimeError("No design satisfies all constraints.")

    costs = np.where(feasible, cost_total(area_CSP, area_PVT), np.inf)
    best = int(np.argmin(costs))
    return int(count_heliostats[best]), int(count_PVT_panels[best])


# =============================================================================
# 6. RUN EVERYTHING FOR ONE REFLECTIVITY
# =============================================================================
def run_case(reflectivity, data, count_heliostats=None, count_PVT_panels=None):
    """
    Builds yields for a given mirror reflectivity. If no design is given, solves for the
    optimal one. Returns everything needed for printing and plotting.
    """
    irradiance, T_ambient, demand_heat, demand_elec = data
    eff_optical, eff_CSP = efficiency_chain(reflectivity)
    yields = yields_per_m2(irradiance, T_ambient, eff_CSP)
    sun_hrs = sun_hrs_mask(irradiance)

    if count_heliostats is None:
        count_heliostats, count_PVT_panels = solve(yields, sun_hrs, demand_heat, demand_elec)

    area_CSP = count_heliostats * area_per_heliostat
    area_PVT = count_PVT_panels * area_per_PVT_panel
    supply_CSP_elec, supply_PVT_elec, supply_PVT_heat = supply(area_CSP, area_PVT, *yields)

    return {
        "reflectivity": reflectivity, "eff_optical": eff_optical, "eff_CSP": eff_CSP,
        "yields": yields, "sun_hrs": sun_hrs,
        "count_heliostats": count_heliostats, "count_PVT_panels": count_PVT_panels,
        "area_CSP": area_CSP, "area_PVT": area_PVT,
        "cost_total": cost_total(area_CSP, area_PVT),
        "supply_CSP_elec": supply_CSP_elec, "supply_PVT_elec": supply_PVT_elec,
        "supply_PVT_heat": supply_PVT_heat,
        "T_water_out": outlet_temperature(yields[2]),
        "satisfied": check_constraints(area_CSP, area_PVT, yields, sun_hrs, demand_heat, demand_elec),
    }


def print_results(result, data):
    """Prints the design, the cost and a constraint check, plus daily energy totals."""
    _, _, demand_heat, demand_elec = data
    print(f"\n=== Reflectivity {result['reflectivity']:.2f} ===")
    print(f"Eff_optical = {result['eff_optical']:.4f}   Eff_CSP = {result['eff_CSP']:.4f}")
    print(f"Count_heliostats = {result['count_heliostats']:,}   (Area_CSP = {result['area_CSP']:,.0f} m2)")
    print(f"Count_PVT_panels = {result['count_PVT_panels']:,}   (Area_PVT = {result['area_PVT']:,.0f} m2)")
    print(f"Area used        = {result['area_CSP'] + result['area_PVT']:,.0f} m2 of {area_max:,.0f} m2")
    print(f"Cost_total       = EUR {result['cost_total']:,.0f}")
    for name, ok in result["satisfied"].items():
        print(f"  {name:<28} {'OK' if ok else 'VIOLATED'}")
    print("  Max PVT outlet temperature: " + ", ".join(
        f"{day} {result['T_water_out'][day].max():.1f} degC" for day in days))
    for day in days:
        sun = result["sun_hrs"][day]
        csp = result["supply_CSP_elec"][day]
        pvt = result["supply_PVT_elec"][day]
        heat = result["supply_PVT_heat"][day]
        print(f"  {day}: sunlit hours {int(sun.sum())}")
        print(f"    electricity  supply {csp[sun].sum() + pvt[sun].sum():7.1f} GJ  vs demand "
              f"{demand_elec[day][sun].sum():7.1f} GJ   (CSP share {csp.sum() / (csp.sum() + pvt.sum()):.1%})")
        print(f"    heat         supply {heat[sun].sum():7.1f} GJ  vs demand {demand_heat[day][sun].sum():7.1f} GJ")
        print(f"    whole day    electricity supply {csp.sum() + pvt.sum():7.1f} / demand {demand_elec[day].sum():7.1f} GJ;"
              f"  heat supply {heat.sum():7.1f} / demand {demand_heat[day].sum():7.1f} GJ")


# =============================================================================
# 7. FIGURES
# =============================================================================
COLOR_CSP = "#2a78d6"       # blue   : CSP electricity
COLOR_PVT = "#eb6834"       # orange : PVT (electricity and heat)
COLOR_ALT = "#1baf7a"       # aqua   : second case in comparisons
COLOR_DEMAND = "#0b0b0b"    # ink    : demand
COLOR_GRID = "#e4e3df"
COLOR_TEXT_2 = "#52514e"

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10,
    "axes.edgecolor": COLOR_TEXT_2, "axes.labelcolor": "#0b0b0b",
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": COLOR_GRID, "grid.linewidth": 0.8,
    "xtick.color": COLOR_TEXT_2, "ytick.color": COLOR_TEXT_2,
    "legend.frameon": False,
})


def figure_cost_map(result, data, path):
    """
    Q5: the design space. Every (PVT panels, heliostats) pair on a grid is coloured by its
    total cost if it satisfies all constraints, grey if it violates one. The boundary lines
    of the binding constraints and the optimum are drawn on top.
    """
    irradiance, T_ambient, demand_heat, demand_elec = data
    yields, sun_hrs = result["yields"], result["sun_hrs"]

    count_PVT_axis = np.linspace(0, 3 * result["count_PVT_panels"], 400)
    count_helio_axis = np.linspace(0, 3 * result["count_heliostats"], 400)
    grid_PVT, grid_helio = np.meshgrid(count_PVT_axis, count_helio_axis)
    area_PVT_grid = grid_PVT * area_per_PVT_panel
    area_CSP_grid = grid_helio * area_per_heliostat

    satisfied = check_constraints(area_CSP_grid, area_PVT_grid, yields, sun_hrs, demand_heat, demand_elec)
    feasible = np.logical_and.reduce(list(satisfied.values()))
    cost_grid_M = np.where(feasible, cost_total(area_CSP_grid, area_PVT_grid) / 1e6, np.nan)

    fig, ax = plt.subplots(figsize=(8, 5.6))
    ax.grid(False)
    ax.contourf(grid_PVT, grid_helio, ~feasible, levels=[0.5, 1.5], colors=["#efeeea"])
    filled = ax.contourf(grid_PVT, grid_helio, cost_grid_M, levels=14, cmap="Blues")
    colorbar = fig.colorbar(filled, ax=ax)
    colorbar.set_label("Total cost (million EUR)")

    # Boundary lines (where each constraint holds with equality), one per constraint type
    boundary_styles = {"C1": ("--", "C1 heat"), "C2": ("-", "C2 electricity"),
                       "C3": (":", "C3 CSP share ≥ 60%"), "C4": ("-.", "C4 area ≤ 450,000 m²")}
    for key, (style, label) in boundary_styles.items():
        drawn = False
        for name, ok in satisfied.items():
            # Skip a constraint whose boundary lies outside the plotted range (e.g. C4 here)
            if not name.startswith(key) or ok.all() or not ok.any():
                continue
            # C3 is the same line on both days (irradiance cancels), so draw it once, unlabelled
            if key == "C3" and drawn:
                continue
            lines = ax.contour(grid_PVT, grid_helio, ok.astype(float), levels=[0.5],
                               colors=COLOR_DEMAND, linestyles=style, linewidths=1.2)
            # Label each line with its day, e.g. "June" / "December"
            day_label = name[name.find("(") + 1:name.find(")")] if "(" in name else ""
            if day_label and key != "C3":
                ax.clabel(lines, fmt={0.5: day_label}, fontsize=7.5, inline=True)
            if not drawn:
                ax.plot([], [], color=COLOR_DEMAND, linestyle=style, linewidth=1.2, label=label)
                drawn = True

    ax.plot(result["count_PVT_panels"], result["count_heliostats"], "o", markersize=9,
            markerfacecolor=COLOR_PVT, markeredgecolor="white", markeredgewidth=2, zorder=5,
            label="Optimum")
    ax.annotate(f"Optimum\n{result['count_PVT_panels']:,} PVT panels\n{result['count_heliostats']:,} heliostats\n"
                f"EUR {result['cost_total'] / 1e6:.2f} M",
                (result["count_PVT_panels"], result["count_heliostats"]),
                xytext=(18, 18), textcoords="offset points", fontsize=9,
                bbox=dict(boxstyle="round,pad=0.3", fc="white", ec=COLOR_GRID))
    ax.set_xlabel("Count of PVT panels (4 m² each)")
    ax.set_ylabel("Count of heliostats (220 m² each)")
    ax.set_title("Design space: total cost of every feasible design (grey = infeasible)\n"
                 "C4 (area ≤ 450,000 m²) lies outside the plotted range and is not binding",
                 loc="left", fontsize=10.5)
    ax.legend(loc="upper right", fontsize=8.5)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def figure_supply_vs_demand(result, data, path):
    """
    Q6: hourly plant output vs demand, both days, electricity (stacked CSP + PVT) and heat.
    Units: GJ per hour (same as the demand data). 1 GJ/h = 0.278 MW.
    """
    _, _, demand_heat, demand_elec = data
    fig, axes = plt.subplots(2, 2, figsize=(12, 7.5), sharex=True)
    width = 0.8
    for col, day in enumerate(days):
        csp = result["supply_CSP_elec"][day]
        pvt_e = result["supply_PVT_elec"][day]
        heat = result["supply_PVT_heat"][day]
        hours = np.array(hrs)

        ax = axes[0, col]
        ax.bar(hours, csp, width, color=COLOR_CSP, label="CSP electricity", edgecolor="white", linewidth=1)
        ax.bar(hours, pvt_e, width, bottom=csp, color=COLOR_PVT, label="PVT electricity",
               edgecolor="white", linewidth=1)
        ax.step(hours, demand_elec[day], where="mid", color=COLOR_DEMAND, linewidth=2, label="Electricity demand")
        ax.set_title(f"Electricity, {day}", loc="left")
        ax.set_ylabel("Energy per hour (GJ)")

        ax = axes[1, col]
        ax.bar(hours, heat, width, color=COLOR_PVT, label="PVT heat", edgecolor="white", linewidth=1)
        ax.step(hours, demand_heat[day], where="mid", color=COLOR_DEMAND, linewidth=2, label="Heat demand")
        ax.set_title(f"Heat, {day}", loc="left")
        ax.set_ylabel("Energy per hour (GJ)")
        ax.set_xlabel("Hour of day")
        ax.set_xticks(range(0, 24, 2))

    axes[0, 0].legend(loc="upper left", fontsize=8.5)
    axes[1, 0].legend(loc="upper left", fontsize=8.5)
    for ax in axes.flat:
        ax.set_axisbelow(True)
    fig.suptitle(f"Hourly plant output vs demand ({result['count_heliostats']:,} heliostats, "
                 f"{result['count_PVT_panels']:,} PVT panels)", x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def figure_reflectivity(result_095, result_085, data, path):
    """
    Q7: same plant (same counts), mirror reflectivity 0.95 vs 0.85.
    Only CSP electricity changes; PVT electricity and heat are unaffected.
    """
    _, _, demand_heat, demand_elec = data
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.2), sharey=True)
    for col, day in enumerate(days):
        ax = axes[col]
        hours = np.array(hrs)
        for res, color, label in [(result_095, COLOR_CSP, "Reflectivity 0.95"),
                                  (result_085, COLOR_ALT, "Reflectivity 0.85")]:
            total = res["supply_CSP_elec"][day] + res["supply_PVT_elec"][day]
            ax.plot(hours, total, color=color, linewidth=2, marker="o", markersize=4, label=f"Total electricity, {label}")
        ax.step(hours, demand_elec[day], where="mid", color=COLOR_DEMAND, linewidth=2, label="Electricity demand")
        ax.set_title(f"Electricity output, {day}", loc="left")
        ax.set_xlabel("Hour of day")
        ax.set_xticks(range(0, 24, 2))
    axes[0].set_ylabel("Energy per hour (GJ)")
    axes[0].legend(loc="upper left", fontsize=8.5)
    fig.suptitle("Effect of mirror degradation on plant output (heat output is unchanged)",
                 x=0.01, ha="left", fontsize=12)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


# =============================================================================
# 8. EXCEL EXPORT: the data behind the figures
# =============================================================================
def export_workbook(result, data, path):
    """
    Writes the results workbook. Every sheet starts with a title row, every column header
    states its quantity and unit in plain words, so the file reads without a separate guide.

    Sheets:
      Inputs             every given value and the chosen design (yellow = design choice)
      June hourly        hour-by-hour data, output per m2, plant supply, demand and balance
      December hourly    same for December
      Cost sweep         data behind the design-space figure (Q5)

    The hourly sheets are live formulas linked to Inputs: change the number of heliostats or
    PVT panels on Inputs and every hourly value updates.
    Colour code: blue = given number, black = formula, green = formula using the Inputs sheet.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    irradiance, T_ambient, demand_heat, demand_elec = data
    FONT = "Arial"
    f_title = Font(name=FONT, bold=True, size=13)
    f_sub = Font(name=FONT, italic=True, color="555555")
    f_head = Font(name=FONT, bold=True, color="FFFFFF")
    f_section = Font(name=FONT, bold=True)
    f_text = Font(name=FONT)
    f_input = Font(name=FONT, color="0000FF")
    f_formula = Font(name=FONT, color="000000")
    f_link = Font(name=FONT, color="008000")
    f_total = Font(name=FONT, bold=True)
    fill_head = PatternFill("solid", fgColor="2E3A48")
    fill_design = PatternFill("solid", fgColor="FFFF00")
    fill_total = PatternFill("solid", fgColor="EEF1F5")
    thin_top = Border(top=Side(style="thin", color="2E3A48"))

    def header_row(ws, r, labels, height=48):
        for c, label in enumerate(labels, start=1):
            cell = ws.cell(row=r, column=c, value=label)
            cell.font, cell.fill = f_head, fill_head
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[r].height = height

    wb = Workbook()

    # ------------------------------------------------------------------ Inputs
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "Inputs: given values and chosen design"
    ws["A1"].font = f_title
    ws["A2"] = "Blue = given in the assignment brief.  Yellow = design choice (editable).  Black = calculated."
    ws["A2"].font = f_sub
    header_row(ws, 4, ["Quantity", "Value", "Unit", "Symbol in the mathematical model"], height=24)
    cell_of = {}  # key -> absolute reference on Inputs

    sections = [
        ("CSP plant (heliostat field, receiver, Rankine cycle)", [
            ("reflectivity", "Mirror reflectivity", eff_reflectivity, "fraction", "Eff_reflectivity"),
            ("cosine", "Average cosine efficiency", eff_cosine, "fraction", "Eff_cosine"),
            ("shading", "Shading and blocking efficiency", eff_shading_blocking, "fraction", "Eff_shading blocking"),
            ("atmos", "Atmospheric attenuation efficiency", eff_atmospheric, "fraction", "Eff_atmospheric"),
            ("spill", "Spillage efficiency", eff_spillage, "fraction", "Eff_spillage"),
            ("receiver", "Receiver absorption efficiency", eff_receiver, "fraction", "Eff_receiver"),
            ("rankine", "Rankine cycle thermal efficiency", eff_rankine, "fraction", "Eff_rankine"),
            ("area_helio", "Mirror area per heliostat", area_per_heliostat, "m²", "Area_per heliostat"),
            ("cost_csp", "Heliostat cost per m²", cost_CSP, "EUR/m²", "Cost_CSP"),
        ]),
        ("PVT panels (flat plate collector)", [
            ("pv_eff", "PVT electrical efficiency", eff_PVT_elec, "fraction", "Eff_PVT,elec"),
            ("absorbed", "Effective optical fraction (share of sunlight absorbed)", frac_absorbed, "fraction", "Frac_absorbed"),
            ("U", "Overall heat loss coefficient", coef_heat_loss, "W/(m²·K)", "Coef_heat loss"),
            ("FR", "Collector heat removal factor", factor_heat_removal, "fraction", "Factor_heat removal"),
            ("Tin", "Inlet water temperature", T_water_in, "°C", "T_water in"),
            ("flow", "Water flow rate through one PVT panel", flow_rate_per_PVT_panel, "kg/s", "FlowRate_per PVT panel"),
            ("cp", "Specific heat of water (standard value)", cp_water, "J/(kg·K)", "c_p,water"),
            ("area_pvt", "Area per PVT panel (2 m × 2 m)", area_per_PVT_panel, "m²", "Area_per PVT panel"),
            ("cost_pvt", "PVT cost per m²", cost_PVT, "EUR/m²", "Cost_PVT"),
        ]),
        ("Site and rules", [
            ("area_max", "Maximum available collector area", area_max, "m²", "Area_max"),
            ("share_min", "Minimum share of electricity from CSP", share_CSP_min, "fraction", "Share_CSP,min"),
            ("reflectivity_q7", "Degraded mirror reflectivity (Q7)", 0.85, "fraction", "–"),
        ]),
        ("Unit conversions", [
            ("gj_per_w_hr", "Energy of 1 W sustained for 1 hour", GJ_per_W_hr, "GJ", "GJ_per W-hr"),
        ]),
    ]
    r = 5
    for title, items in sections:
        ws.cell(row=r, column=1, value=title).font = f_section
        r += 1
        for key, label, value, unit, symbol in items:
            ws.cell(row=r, column=1, value=label).font = f_text
            v = ws.cell(row=r, column=2, value=value)
            v.font = f_input
            ws.cell(row=r, column=3, value=unit).font = f_text
            ws.cell(row=r, column=4, value=symbol).font = f_sub
            cell_of[key] = f"Inputs!$B${r}"
            if key == "sigma":
                v.number_format = "0.00E+00"
            elif key == "gj_per_w_hr":
                v.number_format = "0.0E+00"
            r += 1
        r += 1

    def ref(key):  # reference on the Inputs sheet itself (no sheet prefix)
        return cell_of[key].replace("Inputs!", "")

    ws.cell(row=r, column=1, value="Chosen design (cost optimum)").font = f_section
    r += 1
    for key, label, value, symbol in [
        ("n_helio", "Number of heliostats", result["count_heliostats"], "Count_heliostats"),
        ("n_pvt", "Number of PVT panels", result["count_PVT_panels"], "Count_PVT panels"),
    ]:
        ws.cell(row=r, column=1, value=label).font = f_text
        v = ws.cell(row=r, column=2, value=value)
        v.font, v.fill, v.number_format = f_input, fill_design, "#,##0"
        ws.cell(row=r, column=3, value="units").font = f_text
        ws.cell(row=r, column=4, value=symbol).font = f_sub
        cell_of[key] = f"Inputs!$B${r}"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Calculated from the above").font = f_section
    r += 1
    calculated = [
        ("area_csp_tot", "Total mirror area", f"={ref('n_helio')}*{ref('area_helio')}", "m²", "Area_CSP", "#,##0"),
        ("area_pvt_tot", "Total PVT area", f"={ref('n_pvt')}*{ref('area_pvt')}", "m²", "Area_PVT", "#,##0"),
        ("area_used", "Total collector area used", None, "m²", "", "#,##0"),
        ("eff_opt", "Optical efficiency of the heliostat field",
         f"={ref('reflectivity')}*{ref('cosine')}*{ref('shading')}*{ref('atmos')}*{ref('spill')}",
         "fraction", "Eff_optical", "0.000"),
        ("eff_csp", "CSP efficiency, sunlight to electricity",
         None, "fraction", "Eff_CSP", "0.000"),
        ("eff_csp_q7", "CSP efficiency with degraded mirrors (Q7)",
         None, "fraction", "–", "0.000"),
        ("cost_csp_tot", "Cost of heliostats", None, "EUR", "", "#,##0"),
        ("cost_pvt_tot", "Cost of PVT panels", None, "EUR", "", "#,##0"),
        ("cost_total", "Total investment cost", None, "EUR", "Cost_total", "#,##0"),
    ]
    for key, label, formula, unit, symbol, fmt in calculated:
        cell_of[key] = f"Inputs!$B${r}"
        if key == "area_used":
            formula = f"={ref('area_csp_tot')}+{ref('area_pvt_tot')}"
        elif key == "eff_csp":
            formula = f"={ref('eff_opt')}*{ref('receiver')}*{ref('rankine')}"
        elif key == "eff_csp_q7":
            formula = (f"={ref('reflectivity_q7')}*{ref('cosine')}*{ref('shading')}*{ref('atmos')}"
                       f"*{ref('spill')}*{ref('receiver')}*{ref('rankine')}")
        elif key == "cost_csp_tot":
            formula = f"={ref('cost_csp')}*{ref('area_csp_tot')}"
        elif key == "cost_pvt_tot":
            formula = f"={ref('cost_pvt')}*{ref('area_pvt_tot')}"
        elif key == "cost_total":
            formula = f"={ref('cost_csp_tot')}+{ref('cost_pvt_tot')}"
        ws.cell(row=r, column=1, value=label).font = f_total if key == "cost_total" else f_text
        v = ws.cell(row=r, column=2, value=formula)
        v.font, v.number_format = (f_total if key == "cost_total" else f_formula), fmt
        ws.cell(row=r, column=3, value=unit).font = f_text
        ws.cell(row=r, column=4, value=symbol).font = f_sub
        r += 1
    for col, width in zip("ABCD", (52, 16, 12, 28)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"

    I = cell_of  # shorthand for formulas below

    # ------------------------------------------------------------------ Hourly sheets
    # (header, formula template or None for data; {r} = row number)
    columns = [
        ("Hour of day", None, "0"),
        ("Solar irradiance\n[W/m²]", None, "0"),
        ("Ambient temperature\n[°C]", None, "0.0"),
        ("Electricity demand\n[GJ]", None, "0.00"),
        ("Heat demand\n[GJ]", None, "0.00"),
        ("Sunlit hour\n[1 = yes]", "=IF(B{r}>0,1,0)", "0"),
        ("CSP electricity per m² of mirror\n[W/m²]", f"={I['eff_csp']}*B{{r}}", "0.0"),
        ("PVT electricity per m² of panel\n[W/m²]", f"={I['pv_eff']}*B{{r}}", "0.0"),
        ("PVT: sunlight absorbed\n[W/m²]", f"={I['absorbed']}*B{{r}}", "0.0"),
        ("PVT: convection loss\n[W/m²]", f"={I['U']}*({I['Tin']}-C{{r}})", "0.0"),
        ("PVT useful heat per m² of panel\n[W/m²]", f"=MAX(0,{I['FR']}*(I{{r}}-J{{r}}))", "0.0"),
        ("PVT water outlet temperature\n[°C]",
         f"={I['Tin']}+K{{r}}*{I['area_pvt']}/({I['flow']}*{I['cp']})", "0.00"),
        ("Electricity supplied by CSP\n[GJ]", f"={I['gj_per_w_hr']}*{I['area_csp_tot']}*G{{r}}", "0.00"),
        ("Electricity supplied by PVT\n[GJ]", f"={I['gj_per_w_hr']}*{I['area_pvt_tot']}*H{{r}}", "0.00"),
        ("Total electricity supplied\n[GJ]", "=M{r}+N{r}", "0.00"),
        ("Heat supplied by PVT\n[GJ]", f"={I['gj_per_w_hr']}*{I['area_pvt_tot']}*K{{r}}", "0.00"),
        ("Electricity surplus (+) or shortfall (−)\n[GJ]", "=O{r}-D{r}", "0.00"),
        ("Heat surplus (+) or shortfall (−)\n[GJ]", "=P{r}-E{r}", "0.00"),
        ("Q7: electricity by CSP, degraded mirrors\n[GJ]",
         f"={I['gj_per_w_hr']}*{I['area_csp_tot']}*{I['eff_csp_q7']}*B{{r}}", "0.00"),
        ("Q7: total electricity, degraded mirrors\n[GJ]", "=S{r}+N{r}", "0.00"),
    ]
    first, last = 4, 27  # data rows for hours 0..23
    for day in days:
        ws = wb.create_sheet(f"{day} hourly")
        ws["A1"] = f"{day} 2019: hourly plant output and demand"
        ws["A1"].font = f_title
        ws["A2"] = ("Design from the Inputs sheet. Supply and demand are energy delivered within each hour. "
                    "Per-m² outputs are power.")
        ws["A2"].font = f_sub
        header_row(ws, 3, [h for h, _, _ in columns], height=62)
        for hr in hrs:
            rr = first + hr
            data_values = [hr, float(irradiance[day][hr]), float(T_ambient[day][hr]),
                           float(demand_elec[day][hr]), float(demand_heat[day][hr])]
            for c, (_, template, fmt) in enumerate(columns, start=1):
                if template is None:
                    cell = ws.cell(row=rr, column=c, value=data_values[c - 1])
                    cell.font = f_text if c == 1 else f_input
                else:
                    cell = ws.cell(row=rr, column=c, value=template.format(r=rr))
                    cell.font = f_link if "Inputs!" in template else f_formula
                cell.number_format = fmt
        tr = last + 1
        ws.cell(row=tr, column=1, value="Day total")
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=tr, column=c)
            cell.font, cell.fill, cell.border = f_total, fill_total, thin_top
            if c in (4, 5, 6) or c >= 13:
                L = get_column_letter(c)
                cell.value = f"=SUM({L}{first}:{L}{last})"
                cell.number_format = columns[c - 1][2]
        sr = tr + 1
        ws.cell(row=sr, column=1, value="Sunlit hours only")
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=sr, column=c)
            cell.font, cell.fill = f_total, fill_total
            if c in (4, 5) or c >= 13:
                L = get_column_letter(c)
                cell.value = f"=SUMPRODUCT({L}{first}:{L}{last},$F${first}:$F${last})"
                cell.number_format = columns[c - 1][2]
        ws.column_dimensions["A"].width = 16
        for c in range(2, len(columns) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 15
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ Cost sweep
    ws = wb.create_sheet("Cost sweep")
    ws["A1"] = "Cost sweep: cheapest number of heliostats for each number of PVT panels"
    ws["A1"].font = f_title
    ws["A2"] = ("Data behind the design-space figure. Values calculated by the script. "
                "The highlighted row is the cost optimum.")
    ws["A2"].font = f_sub
    header_row(ws, 3, ["Number of PVT panels",
                       "Fewest heliostats meeting electricity demand and 60% CSP share (both days)",
                       "Heat demand met on both days", "Design feasible", "Total cost\n[EUR]"], height=62)
    yields, sun_hrs = result["yields"], result["sun_hrs"]
    step = 100
    pvt_counts = sorted(set(range(0, 3 * result["count_PVT_panels"] + step, step)) | {result["count_PVT_panels"]})
    for n_pvt in pvt_counts:
        a_pvt = n_pvt * area_per_PVT_panel
        need = 0.0
        for day in days:
            sun = sun_hrs[day]
            e_csp = GJ_per_W_hr * yields[0][day][sun].sum()
            e_pv = GJ_per_W_hr * yields[1][day][sun].sum()
            need = max(need, (demand_elec[day][sun].sum() - a_pvt * e_pv) / e_csp,
                       share_CSP_min / (1 - share_CSP_min) * e_pv / e_csp * a_pvt)
        n_helio = math.ceil(need / area_per_heliostat - 1e-9)
        ok = check_constraints(n_helio * area_per_heliostat, a_pvt, yields, sun_hrs, demand_heat, demand_elec)
        heat_ok = all(v for k, v in ok.items() if k.startswith("C1"))
        ws.append([n_pvt, n_helio, "yes" if heat_ok else "no", "yes" if all(ok.values()) else "no",
                   cost_total(n_helio * area_per_heliostat, a_pvt)])
        rr = ws.max_row
        for c in range(1, 6):
            cell = ws.cell(row=rr, column=c)
            cell.font = f_text
            if n_pvt == result["count_PVT_panels"]:
                cell.fill, cell.font = fill_design, f_total
        ws.cell(row=rr, column=1).number_format = "#,##0"
        ws.cell(row=rr, column=5).number_format = "#,##0"
    for col, width in zip("ABCDE", (16, 30, 16, 12, 16)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"

    wb.save(path)


# =============================================================================
# MAIN
# =============================================================================
if __name__ == "__main__":
    from pathlib import Path

    # Folder layout:  <project>/3_Code (this script), 4_Visualizations, 5_Data
    project_dir = Path(__file__).resolve().parent.parent
    folder_figures = project_dir / "4_Visualizations"
    folder_data = project_dir / "5_Data"
    folder_figures.mkdir(exist_ok=True)
    folder_data.mkdir(exist_ok=True)

    excel_path = Path(sys.argv[1]) if len(sys.argv) > 1 else folder_data / "Input_Data_Assignment_1.xlsx"
    data = read_data(excel_path)

    # Q5: optimal design with the given reflectivity (0.95)
    result_095 = run_case(eff_reflectivity, data)
    print_results(result_095, data)

    # Q7: the SAME plant after mirror degradation to 0.85 (counts fixed, not re-optimised)
    result_085 = run_case(0.85, data, result_095["count_heliostats"], result_095["count_PVT_panels"])
    print_results(result_085, data)

    figure_cost_map(result_095, data, folder_figures / "Q5_Design_Space_Cost_Map.png")
    figure_supply_vs_demand(result_095, data, folder_figures / "Q6_Hourly_Supply_vs_Demand.png")
    figure_reflectivity(result_095, result_085, data, folder_figures / "Q7_Mirror_Degradation.png")
    export_workbook(result_095, data, folder_data / "Results_Plant_Sizing.xlsx")
    print(f"\nSaved figures to {folder_figures} and results workbook to {folder_data}")
